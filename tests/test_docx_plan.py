#!/usr/bin/env python3
"""Plan and apply tests for the DOCX handler."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage
from PIL import Image
# skia-python does not publish PEP 561 stubs; this is the native rendering boundary.
import skia  # type: ignore[import-not-found]

from folder_replacement_test_support import FONT_PATH, _EmptyOcrProvider
from pipeline.folder_replacement.docx import (
    apply_docx,
    docx_work_item_total,
    plan_docx,
    replace_docx_file,
)
from pipeline.folder_replacement.planning import PhaseProgress, plan_and_apply
from pipeline.ocr import BoundingPolygon, OcrRequest, OcrResult, OcrText, PixelPoint
from pipeline.text_replacement import (
    BatchLimits,
    TextReplacementItem,
    TextReplacementRequest,
    TextReplacementResult,
    TranslationEngine,
    replace_each,
)

LAYOUTS = (
    "preserve-source-formatting",
    "preserve-basic-layout",
    "preserve-basic-layout-source-font",
)
_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_CHART_NS = (
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)
_RELATIONSHIPS = "http://schemas.openxmlformats.org/package/2006/relationships"
_RELATIONSHIP_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DOCUMENT = f"""<?xml version="1.0" encoding="UTF-8"?>
<w:document {_W_NS}
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:body>
    <w:p><w:r><w:t>Body text</w:t></w:r></w:p>
    <w:p><w:r><w:drawing><wp:inline>
      <wp:extent cx="914400" cy="457200"/><wp:docPr id="1" name="Box"/>
      <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <wps:wsp><wps:txbx><w:txbxContent>
          <w:p><w:r><w:rPr><w:sz w:val="48"/></w:rPr><w:t>Box text</w:t></w:r></w:p>
        </w:txbxContent></wps:txbx></wps:wsp>
      </a:graphicData></a:graphic>
    </wp:inline></w:drawing></w:r></w:p>
    <w:sectPr><w:pgSz w:w="12240" w:h="15840"/></w:sectPr>
  </w:body>
</w:document>""".encode()


def _word_part(root: str, text: str) -> bytes:
    return (
        f'<?xml version="1.0" encoding="UTF-8"?><w:{root} {_W_NS}>'
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:{root}>"
    ).encode()


class _BatchProvider:
    """Translate by suffixing, and record the texts of each request."""

    def __init__(self, limits: BatchLimits | None = None) -> None:
        self.batch_limits = limits
        self.requests: list[tuple[str, ...]] = []

    def replace(self, request: TextReplacementRequest) -> TextReplacementResult:
        self.requests.append(request.texts)
        return replace_each(request, lambda text: TextReplacementItem(f"{text} translated", 1.0))


class _RegionOcr(_EmptyOcrProvider):
    """Return one confident region for any image."""

    def __init__(self, text: str) -> None:
        self.text = text

    def recognize(self, request: OcrRequest) -> OcrResult:
        polygon = BoundingPolygon(
            (PixelPoint(5, 5), PixelPoint(95, 5), PixelPoint(95, 30), PixelPoint(5, 30))
        )
        return OcrResult((OcrText(self.text, 0.99, polygon),))


def _typeface() -> skia.Typeface:
    typeface = skia.Typeface.MakeFromFile(str(FONT_PATH))
    if typeface is None:
        raise AssertionError("Could not load test typeface.")
    return typeface


def _png() -> bytes:
    output = BytesIO()
    Image.new("RGB", (100, 40), "white").save(output, "PNG")
    return output.getvalue()


def _workbook_bytes(*, image: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet["A1"] = "Cell text"
    if image:
        sheet.add_image(WorkbookImage(BytesIO(_png())), "H8")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _docx_bytes(*, chart: bool = False, image: bool = False, workbook_image: bool = False) -> bytes:
    """Build a synthetic package holding each Word text class; optionally a chart and an image."""
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            f'<?xml version="1.0"?><Relationships xmlns="{_RELATIONSHIPS}"><Relationship Id="rId1" '
            f'Type="{_RELATIONSHIP_TYPE}/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        archive.writestr("word/document.xml", _DOCUMENT)
        archive.writestr(
            "word/_rels/document.xml.rels",
            f'<?xml version="1.0"?><Relationships xmlns="{_RELATIONSHIPS}">'
            + (
                f'<Relationship Id="rIdChart" Type="{_RELATIONSHIP_TYPE}/chart" Target="charts/chart1.xml"/>'
                if chart
                else ""
            )
            + "</Relationships>",
        )
        archive.writestr("word/header1.xml", _word_part("hdr", "Header text"))
        archive.writestr("word/footer1.xml", _word_part("ftr", "Footer text"))
        archive.writestr("word/footnotes.xml", _word_part("footnotes", "Footnote text"))
        archive.writestr("word/endnotes.xml", _word_part("endnotes", "Endnote text"))
        archive.writestr("word/comments.xml", _word_part("comments", "Comment text"))
        if chart:
            _add_chart(archive, workbook_image)
        if image:
            archive.writestr("word/media/image1.png", _png())
    return output.getvalue()


def _add_chart(archive: ZipFile, workbook_image: bool) -> None:
    archive.writestr(
        "word/charts/chart1.xml",
        f'<?xml version="1.0"?><c:chartSpace {_CHART_NS}><c:chart><c:title><c:tx><c:rich>'
        "<a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Chart title</a:t></a:r></a:p></c:rich></c:tx></c:title>"
        '</c:chart><c:externalData r:id="rIdWorkbook"/></c:chartSpace>',
    )
    archive.writestr(
        "word/charts/_rels/chart1.xml.rels",
        f'<?xml version="1.0"?><Relationships xmlns="{_RELATIONSHIPS}">'
        f'<Relationship Id="rIdWorkbook" Type="{_RELATIONSHIP_TYPE}/package" Target="../embeddings/chart.xlsx"/>'
        f'<Relationship Id="rIdShape" Type="{_RELATIONSHIP_TYPE}/chartUserShapes" Target="userShapes/drawing1.xml"/>'
        "</Relationships>",
    )
    archive.writestr(
        "word/charts/userShapes/drawing1.xml",
        '<?xml version="1.0"?><cdr:userShapes xmlns:cdr="http://schemas.openxmlformats.org/drawingml/2006/chartDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><cdr:relSizeAnchor><cdr:sp><cdr:txBody>'
        "<a:p><a:r><a:t>Shape label</a:t></a:r></a:p></cdr:txBody></cdr:sp></cdr:relSizeAnchor></cdr:userShapes>",
    )
    archive.writestr("word/embeddings/chart.xlsx", _workbook_bytes(image=workbook_image))


def _parts(data: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _run_batched(
    data: bytes, layout: str, ocr: _EmptyOcrProvider | None = None,
) -> tuple[bytes, TranslationEngine, _BatchProvider]:
    provider = _BatchProvider()
    plan = plan_docx(data, ocr, "en", "fr", document_text_layout=layout)
    engine = TranslationEngine(provider, "en", "fr")
    result = plan_and_apply(
        plan, engine, lambda translate: apply_docx(plan, translate, _typeface())
    )
    return result.data, engine, provider


class DocxPlanTests(unittest.TestCase):
    # Verifies FR-2026-10-22-04's coverage of every Word text class in every layout mode.
    def test_plan_lists_every_word_text_class_in_every_layout(self) -> None:
        for layout in LAYOUTS:
            with self.subTest(layout=layout):
                plan = plan_docx(_docx_bytes(chart=True), None, "en", "fr", document_text_layout=layout)

                for expected in (
                    "Body text", "Box text", "Header text", "Footer text", "Footnote text",
                    "Endnote text", "Comment text", "Chart title", "Shape label",
                ):
                    self.assertIn(expected, plan.texts)

    # Verifies FR-2026-10-22-04's plan-miss rule and its identical-to-unbatched rule.
    def test_apply_requests_only_planned_text_and_matches_unbatched_output(self) -> None:
        data = _docx_bytes(chart=True, image=True)
        for layout in LAYOUTS:
            with self.subTest(layout=layout), TemporaryDirectory() as directory:
                source = Path(directory) / "source.docx"
                source.write_bytes(data)
                ocr = _RegionOcr("Image words")

                batched, engine, provider = _run_batched(data, layout, ocr)
                unbatched_provider = _BatchProvider(
                    BatchLimits(max_items=1, target_characters=1, max_characters=1000)
                )
                destination = Path(directory) / "unbatched.docx"
                replace_docx_file(
                    source, destination, ocr, unbatched_provider, "en", "fr", _typeface(),
                    lambda _label: None, layout,
                )

                units_with_texts = sum(
                    bool(unit.texts)
                    for unit in plan_docx(data, ocr, "en", "fr", document_text_layout=layout).units()
                )
                self.assertEqual(0, engine.plan_misses)
                self.assertEqual(units_with_texts, len(provider.requests))
                self.assertGreater(len(unbatched_provider.requests), 1)
                self.assertEqual(_parts(destination.read_bytes()), _parts(batched))

    # Verifies FR-2026-10-22-04's rule that translated text lands in the planned parts.
    def test_translated_text_is_written_to_each_part(self) -> None:
        data, _engine, _provider = _run_batched(_docx_bytes(chart=True), LAYOUTS[0])

        parts = _parts(data)
        for name, text in (
            ("word/document.xml", b"Body text translated"),
            ("word/header1.xml", b"Header text translated"),
            ("word/comments.xml", b"Comment text translated"),
            ("word/charts/chart1.xml", b"Chart title translated"),
            ("word/charts/userShapes/drawing1.xml", b"Shape label translated"),
        ):
            self.assertIn(text, parts[name], name)

    # Verifies FR-2026-10-22-05's path-style identifiers for images, charts and nested images.
    def test_nested_documents_have_path_style_unit_identifiers(self) -> None:
        data = _docx_bytes(chart=True, image=True, workbook_image=True)

        plan = plan_docx(data, _RegionOcr("Image words"), "en", "fr")

        self.assertEqual(
            ["docx", "docx/image-1", "docx/chart-1", "docx/chart-1/image-1"],
            [unit.unit_id for unit in plan.units()],
        )
        batched, engine, _provider = _run_batched(data, LAYOUTS[0], _RegionOcr("Image words"))
        self.assertEqual(0, engine.plan_misses)
        self.assertTrue(batched)

    # Verifies FR-2026-10-22-04's one-request-per-unit packing for a document and its nested workbook.
    def test_the_document_and_its_chart_workbook_each_fit_one_batch(self) -> None:
        _data, _engine, provider = _run_batched(_docx_bytes(chart=True), LAYOUTS[1])

        self.assertEqual(2, len(provider.requests))
        self.assertIn("Body text", provider.requests[0])
        self.assertIn("Cell text", provider.requests[1])
        for request in provider.requests:
            self.assertEqual(len(set(request)), len(request))

    # Verifies FR-2026-10-22-07's plan and apply items, which come from the part names alone.
    def test_work_item_total_matches_the_plan_and_apply_ticks(self) -> None:
        data = _docx_bytes(chart=True, image=True)
        labels: list[str] = []
        progress = PhaseProgress(labels.append, labels.append, lambda _extra: None)
        plan = plan_docx(data, _RegionOcr("Image words"), "en", "fr", progress=progress)

        plan_and_apply(
            plan,
            TranslationEngine(_BatchProvider(), "en", "fr"),
            lambda translate: apply_docx(plan, translate, _typeface(), progress=progress),
            progress,
        )
        total = docx_work_item_total(_parts(data))

        self.assertEqual(total, sum(label.startswith("Planning ") for label in labels))
        self.assertEqual(total, sum(label.startswith("Applying ") for label in labels))
        self.assertTrue(labels[0].endswith(f"1 of {total}"))


if __name__ == "__main__":
    unittest.main()
