"""Synthetic tests that one unsupported Office text container never fails the document."""

from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from pipeline.bounded_text_layout import noto_typefaces
from pipeline.folder_replacement.xlsx import _replace_drawing
from folder_replacement_test_support import (
    FolderReplacementTestCase,
    _EmptyOcrProvider,
    _MappedReplacementProvider,
    _RecordingReplacementProvider,
)

# No approved portable face covers a private-use code point.
_UNCOVERED = "\ue000"


class OfficeContainerRecoveryTests(FolderReplacementTestCase):
    def _unsupported_entries(self, sidecar: Path) -> list[dict[str, object]]:
        entries = json.loads(sidecar.read_text(encoding="utf-8"))["entries"]
        return [entry for entry in entries if entry.get("kind") == "unsupported"]

    # Verifies FR-2026-08-27-09.
    def test_pptx_retains_an_unsupported_frame_and_table_cell_and_replaces_their_neighbours(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_root = root / "input"
            output_root = root / "output"
            input_root.mkdir()
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Old"
            slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text = "New"
            table = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
            table.cell(0, 0).text = "Old cell"
            table.cell(0, 1).text = "New cell"
            presentation.save(str(input_root / "deck.pptx"))
            provider = _MappedReplacementProvider({
                "Old": _UNCOVERED, "New": "Fine",
                "Old cell": _UNCOVERED, "New cell": "Fine cell",
            })

            result = self._run(
                input_root, output_root, _EmptyOcrProvider(), provider,
                document_text_layout="preserve-basic-layout", diagnostics_enabled=True,
            )

            self.assertEqual(1, result.processed_files)
            self.assertEqual(0, result.failed_files)
            output = Presentation(str(output_root / "deck.pptx"))
            texts = [shape.text_frame.text for shape in output.slides[0].shapes if shape.has_text_frame]
            self.assertEqual(["Old", "Fine"], texts)
            output_table = next(shape.table for shape in output.slides[0].shapes if shape.has_table)
            self.assertEqual("Old cell", output_table.cell(0, 0).text)
            self.assertEqual("Fine cell", output_table.cell(0, 1).text)
            entries = self._unsupported_entries(output_root / "deck.pptx.diagnostics.json")
            self.assertEqual(
                ["pptx_table_cell", "pptx_text_frame"],
                sorted(str(entry["container_kind"]) for entry in entries),
            )
            for entry in entries:
                self.assertEqual(_UNCOVERED, entry["replacement_text"])
                self.assertEqual(["U+E000"], entry["code_points"])
            self.assertEqual(1, [request.text for request in provider.requests].count("Old"))

    # Verifies FR-2026-08-27-09.
    def test_docx_retains_an_unsupported_text_box_and_replaces_flowing_text(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_root = root / "input"
            output_root = root / "output"
            input_root.mkdir()
            self._write_complete_docx(input_root / "document.docx")
            provider = _MappedReplacementProvider({
                "Flow text": "Fine", "Bold text": _UNCOVERED, "Italic text": "Fine",
            })

            result = self._run(
                input_root, output_root, _EmptyOcrProvider(), provider,
                document_text_layout="preserve-basic-layout", diagnostics_enabled=True,
            )

            self.assertEqual(1, result.processed_files)
            self.assertEqual(0, result.failed_files)
            with ZipFile(output_root / "document.docx") as archive:
                document = archive.read("word/document.xml").decode("utf-8")
            ElementTree.fromstring(document)
            self.assertNotIn("Flow text", document)
            for text in ("Bold text", "Italic text", "Underlined text"):
                self.assertIn(text, document)
            self.assertNotIn(_UNCOVERED, document)
            entries = self._unsupported_entries(output_root / "document.docx.diagnostics.json")
            self.assertEqual(["docx_textbox"], [str(entry["container_kind"]) for entry in entries])
            self.assertEqual(["U+E000"], entries[0]["code_points"])

    # Verifies FR-2026-08-27-09.
    def test_xlsx_retains_an_unsupported_bounded_cell_and_replaces_its_neighbour(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_root = root / "input"
            output_root = root / "output"
            input_root.mkdir()
            workbook = Workbook()
            worksheet = workbook.active
            assert worksheet is not None
            worksheet.column_dimensions["A"].width = 12
            worksheet.row_dimensions[1].height = 30
            worksheet.row_dimensions[2].height = 30
            worksheet["A1"] = "Old"
            worksheet["A2"] = "New"
            workbook.save(input_root / "workbook.xlsx")
            provider = _MappedReplacementProvider({"Old": _UNCOVERED, "New": "Fine"})

            result = self._run(
                input_root, output_root, _EmptyOcrProvider(), provider,
                document_text_layout="preserve-basic-layout", diagnostics_enabled=True,
            )

            self.assertEqual(1, result.processed_files)
            self.assertEqual(0, result.failed_files)
            output_sheet = load_workbook(output_root / "workbook.xlsx").active
            assert output_sheet is not None
            self.assertEqual("Old", output_sheet["A1"].value)
            self.assertEqual("Fine", output_sheet["A2"].value)
            entries = self._unsupported_entries(output_root / "workbook.xlsx.diagnostics.json")
            self.assertEqual(["xlsx_cell"], [str(entry["container_kind"]) for entry in entries])
            self.assertEqual("A1", entries[0]["cell"])
            self.assertEqual(["U+E000"], entries[0]["code_points"])

    # Verifies FR-2026-08-27-09.
    def test_xlsx_drawing_retains_an_unsupported_shape_without_translating_it_unfitted(self) -> None:
        def shape(text: str) -> str:
            return (
                '<xdr:sp><xdr:spPr><a:xfrm><a:ext cx="914400" cy="457200"/></a:xfrm></xdr:spPr>'
                '<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr sz="1200"/>'
                f"<a:t>{text}</a:t></a:r></a:p></xdr:txBody></xdr:sp>"
            )

        drawing = (
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            + shape("Old") + shape("New") + "</xdr:wsDr>"
        ).encode()
        provider = _MappedReplacementProvider({"Old": _UNCOVERED, "New": "Fine"})
        diagnostics: list[dict[str, object]] = []

        updated, count = _replace_drawing(
            drawing, provider, "en", "en", noto_typefaces(), True,
            part_name="xl/drawings/drawing1.xml", diagnostics=diagnostics,
        )

        text = updated.decode("utf-8")
        ElementTree.fromstring(text)
        self.assertIn(">Old<", text)
        self.assertIn(">Fine<", text)
        self.assertNotIn(_UNCOVERED, text)
        self.assertEqual(1, count)
        self.assertEqual(["xlsx_drawing_shape"], [str(entry["container_kind"]) for entry in diagnostics])
        self.assertEqual("xl/drawings/drawing1.xml", diagnostics[0]["package_part"])
        self.assertEqual(_UNCOVERED, diagnostics[0]["replacement_text"])
        self.assertEqual(1, [request.text for request in provider.requests].count("Old"))

    # Verifies FR-2026-10-23-01.
    def test_pptx_text_frame_with_a_tab_is_fitted_and_keeps_its_tab(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_root = root / "input"
            output_root = root / "output"
            input_root.mkdir()
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Source"
            presentation.save(str(input_root / "deck.pptx"))

            result = self._run(
                input_root, output_root, _EmptyOcrProvider(),
                _RecordingReplacementProvider(replacement_text="Left\tRight"),
                document_text_layout="preserve-basic-layout", diagnostics_enabled=True,
            )

            self.assertEqual(1, result.processed_files)
            self.assertEqual(0, result.failed_files)
            output = Presentation(str(output_root / "deck.pptx"))
            frame = next(shape for shape in output.slides[0].shapes if shape.has_text_frame)
            self.assertIn("Left", frame.text_frame.text)
            self.assertIn("Right", frame.text_frame.text)
            self.assertEqual([], self._unsupported_entries(output_root / "deck.pptx.diagnostics.json")
                             if (output_root / "deck.pptx.diagnostics.json").exists() else [])
