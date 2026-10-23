#!/usr/bin/env python3
"""Plan and apply tests for the PPTX handler."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
# skia-python does not publish PEP 561 stubs; this is the native rendering boundary.
import skia  # type: ignore[import-not-found]

from folder_replacement_test_support import FONT_PATH, _EmptyOcrProvider
from pipeline.folder_replacement.planning import PhaseProgress, plan_and_apply
from pipeline.folder_replacement.pptx import (
    apply_pptx,
    plan_pptx,
    pptx_work_item_total,
    replace_pptx_file,
)
from pipeline.ocr import BoundingPolygon, OcrRequest, OcrResult, OcrText, PixelPoint
from pipeline.text_replacement import (
    BatchLimits,
    TextReplacementItem,
    TextReplacementRequest,
    TextReplacementResult,
    TranslationEngine,
    replace_each,
)

LAYOUTS = (
    "preserve-source-formatting",
    "preserve-basic-layout",
    "preserve-basic-layout-source-font",
)
_SMARTART_DATA = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <dgm:ptLst>
    <dgm:pt modelId="node-1"><dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>SmartArt first</a:t></a:r></a:p></dgm:t></dgm:pt>
  </dgm:ptLst>
  <dgm:cxnLst/><dgm:bg/><dgm:whole/>
</dgm:dataModel>"""


class _BatchProvider:
    """Translate by suffixing, and record the texts of each request."""

    def __init__(self, limits: BatchLimits | None = None) -> None:
        self.batch_limits = limits
        self.requests: list[tuple[str, ...]] = []

    def replace(self, request: TextReplacementRequest) -> TextReplacementResult:
        self.requests.append(request.texts)
        return replace_each(request, lambda text: TextReplacementItem(f"{text} translated", 1.0))


class _RegionOcr(_EmptyOcrProvider):
    """Return one confident region for any image."""

    def __init__(self, text: str) -> None:
        self.text = text

    def recognize(self, request: OcrRequest) -> OcrResult:
        polygon = BoundingPolygon(
            (PixelPoint(5, 5), PixelPoint(95, 5), PixelPoint(95, 30), PixelPoint(5, 30))
        )
        return OcrResult((OcrText(self.text, 0.99, polygon),))


def _typeface() -> skia.Typeface:
    typeface = skia.Typeface.MakeFromFile(str(FONT_PATH))
    if typeface is None:
        raise AssertionError("Could not load test typeface.")
    return typeface


def _png() -> bytes:
    output = BytesIO()
    Image.new("RGB", (100, 40), "white").save(output, "PNG")
    return output.getvalue()


def _pptx_bytes(*, image: bool = False) -> bytes:
    """Build a synthetic deck holding each slide text class, and optionally an image."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Slide text"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2.5), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Cell one"
    table.cell(0, 1).text = "Cell two"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1)).text_frame.text = "Group text"
    slide.notes_slide.notes_text_frame.text = "Speaker note"
    if image:
        slide.shapes.add_picture(BytesIO(_png()), Inches(6), Inches(1))
    output = BytesIO()
    presentation.save(output)
    return _with_extra_parts(output.getvalue())


def _text_shape(text: str) -> bytes:
    return (
        b'<p:sp><p:nvSpPr><p:cNvPr id="99" name="Text"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        b'<p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:t>' + text.encode() + b"</a:t></a:r></a:p>"
        b"</p:txBody></p:sp>"
    )


_CHART = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart><c:title><c:tx>'
    b"<c:rich><a:bodyPr/><a:p><a:r><a:t>Chart title</a:t></a:r></a:p></c:rich></c:tx></c:title>"
    b"</c:chart></c:chartSpace>"
)
_COMMENT = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<p188:cmLst xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main" '
    b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p188:cm id="1">'
    b"<p188:txBody><a:bodyPr/><a:p><a:r><a:t>Comment text</a:t></a:r></a:p></p188:txBody>"
    b"</p188:cm></p188:cmLst>"
)
_RELATIONSHIP_BASE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _with_extra_parts(data: bytes) -> bytes:
    """Add SmartArt, a chart and a comment linked from the first slide, and text in the
    layout and master."""
    new_parts = {
        "ppt/diagrams/data1.xml": _SMARTART_DATA,
        "ppt/charts/chart1.xml": _CHART,
        "ppt/comments/modernComment_1.xml": _COMMENT,
    }
    links = (
        ("rIdSmartArtData", f"{_RELATIONSHIP_BASE}/diagramData", "../diagrams/data1.xml"),
        ("rIdChart", f"{_RELATIONSHIP_BASE}/chart", "../charts/chart1.xml"),
        (
            "rIdComment",
            "http://schemas.microsoft.com/office/2018/10/relationships/comments",
            "../comments/modernComment_1.xml",
        ),
    )
    relationships = "".join(
        f'<Relationship Id="{identifier}" Type="{kind}" Target="{target}"/>'
        for identifier, kind, target in links
    ).encode()
    with ZipFile(BytesIO(data)) as archive:
        entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for entry, payload in entries:
            if entry.filename == "ppt/slides/_rels/slide1.xml.rels":
                payload = payload.replace(b"</Relationships>", relationships + b"</Relationships>")
            elif entry.filename == "ppt/slideLayouts/slideLayout7.xml":
                payload = payload.replace(b"</p:spTree>", _text_shape("Layout text") + b"</p:spTree>")
            elif entry.filename == "ppt/slideMasters/slideMaster1.xml":
                payload = payload.replace(b"</p:spTree>", _text_shape("Master text") + b"</p:spTree>")
            archive.writestr(entry, payload)
        for name, payload in new_parts.items():
            archive.writestr(name, payload)
    return output.getvalue()



def _parts(data: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _run_batched(
    data: bytes, layout: str, ocr: _EmptyOcrProvider | None = None,
) -> tuple[bytes, TranslationEngine, _BatchProvider]:
    provider = _BatchProvider()
    plan = plan_pptx(data, ocr, "en", "fr", document_text_layout=layout)
    engine = TranslationEngine(provider, "en", "fr")
    result = plan_and_apply(
        plan, engine, lambda translate: apply_pptx(plan, translate, _typeface())
    )
    return result.data, engine, provider


class PptxPlanTests(unittest.TestCase):
    # Verifies FR-2026-10-22-04's coverage of every PowerPoint text class in every layout mode.
    def test_plan_lists_every_slide_text_class_in_every_layout(self) -> None:
        for layout in LAYOUTS:
            with self.subTest(layout=layout):
                plan = plan_pptx(_pptx_bytes(), None, "en", "fr", document_text_layout=layout)

                for expected in (
                    "Slide text", "Cell one", "Cell two", "Group text", "Speaker note", "SmartArt first",
                    "Layout text", "Master text", "Chart title", "Comment text",
                ):
                    self.assertIn(expected, plan.texts)

    # Verifies FR-2026-10-22-04's plan-miss rule and its identical-to-unbatched rule.
    def test_apply_requests_only_planned_text_and_matches_unbatched_output(self) -> None:
        data = _pptx_bytes(image=True)
        for layout in LAYOUTS:
            with self.subTest(layout=layout), TemporaryDirectory() as directory:
                source = Path(directory) / "source.pptx"
                source.write_bytes(data)
                ocr = _RegionOcr("Image words")

                batched, engine, provider = _run_batched(data, layout, ocr)
                unbatched_provider = _BatchProvider(
                    BatchLimits(max_items=1, target_characters=1, max_characters=1000)
                )
                destination = Path(directory) / "unbatched.pptx"
                replace_pptx_file(
                    source, destination, ocr, unbatched_provider, "en", "fr", _typeface(),
                    lambda _label: None, layout,
                )

                units_with_texts = sum(
                    bool(unit.texts)
                    for unit in plan_pptx(data, ocr, "en", "fr", document_text_layout=layout).units()
                )
                self.assertEqual(0, engine.plan_misses)
                self.assertEqual(units_with_texts, len(provider.requests))
                self.assertGreater(len(unbatched_provider.requests), 1)
                self.assertEqual(_parts(destination.read_bytes()), _parts(batched))

    # Verifies FR-2026-10-22-04's rule that translated text lands in the planned parts.
    def test_translated_text_is_written_to_each_part(self) -> None:
        for layout in LAYOUTS:
            with self.subTest(layout=layout):
                data, _engine, _provider = _run_batched(_pptx_bytes(), layout)

                parts = _parts(data)
                for name, text in (
                    ("ppt/slides/slide1.xml", b"Slide text translated"),
                    ("ppt/slides/slide1.xml", b"Cell one translated"),
                    ("ppt/slides/slide1.xml", b"Group text translated"),
                    ("ppt/notesSlides/notesSlide1.xml", b"Speaker note translated"),
                    ("ppt/diagrams/data1.xml", b"SmartArt first translated"),
                    ("ppt/slideLayouts/slideLayout7.xml", b"Layout text translated"),
                    ("ppt/slideMasters/slideMaster1.xml", b"Master text translated"),
                    ("ppt/charts/chart1.xml", b"Chart title translated"),
                    ("ppt/comments/modernComment_1.xml", b"Comment text translated"),
                ):
                    self.assertIn(text, parts[name], f"{layout}: {name}")

    # Verifies FR-2026-10-22-08's rule that no text is translated by two PowerPoint routes.
    def test_no_text_is_translated_twice_in_any_layout(self) -> None:
        for layout in LAYOUTS:
            with self.subTest(layout=layout):
                data, engine, _provider = _run_batched(_pptx_bytes(), layout)

                self.assertEqual(0, engine.plan_misses)
                for name, payload in _parts(data).items():
                    self.assertNotIn(b"translated translated", payload, f"{layout}: {name}")

    # Verifies FR-2026-10-22-08's rule that non-text XML and relationships are retained.
    def test_fitted_layouts_keep_the_generic_parts_valid_and_linked(self) -> None:
        for layout in LAYOUTS[1:]:
            with self.subTest(layout=layout):
                data, _engine, _provider = _run_batched(_pptx_bytes(), layout)

                parts = _parts(data)
                self.assertIn(b"../charts/chart1.xml", parts["ppt/slides/_rels/slide1.xml.rels"])
                self.assertIn(b"../comments/modernComment_1.xml", parts["ppt/slides/_rels/slide1.xml.rels"])
                self.assertIn(b"chartSpace", parts["ppt/charts/chart1.xml"])
                Presentation(BytesIO(data))

    # Verifies FR-2026-10-22-05's path-style identifiers for the images of a presentation.
    def test_embedded_images_have_path_style_unit_identifiers(self) -> None:
        plan = plan_pptx(_pptx_bytes(image=True), _RegionOcr("Image words"), "en", "fr")

        self.assertEqual(["pptx", "pptx/image-1"], [unit.unit_id for unit in plan.units()])

    # Verifies FR-2026-10-22-04's one-request-per-unit packing for a presentation and its image.
    def test_the_presentation_and_its_image_each_fit_one_batch(self) -> None:
        _data, _engine, provider = _run_batched(
            _pptx_bytes(image=True), LAYOUTS[1], _RegionOcr("Image words")
        )

        self.assertEqual(2, len(provider.requests))
        self.assertIn("Slide text", provider.requests[0])
        self.assertEqual(["Image words"], list(provider.requests[1]))
        for request in provider.requests:
            self.assertEqual(len(set(request)), len(request))

    # Verifies FR-2026-10-22-07's plan and apply items, which come from the part names alone.
    def test_work_item_total_matches_the_plan_and_apply_ticks(self) -> None:
        for layout in LAYOUTS:
            with self.subTest(layout=layout):
                data = _pptx_bytes(image=True)
                labels: list[str] = []
                progress = PhaseProgress(labels.append, labels.append, lambda _extra: None)
                plan = plan_pptx(
                    data, _RegionOcr("Image words"), "en", "fr",
                    document_text_layout=layout, progress=progress,
                )

                plan_and_apply(
                    plan,
                    TranslationEngine(_BatchProvider(), "en", "fr"),
                    lambda translate: apply_pptx(plan, translate, _typeface(), progress=progress),
                    progress,
                )
                total = pptx_work_item_total(_parts(data))

                self.assertEqual(total, sum(label.startswith("Planning ") for label in labels))
                self.assertEqual(total, sum(label.startswith("Applying ") for label in labels))
                self.assertTrue(labels[0].endswith(f"1 of {total}"))

    # Verifies FR-2026-10-22-04's rejection of an unknown layout before any text is requested.
    def test_plan_rejects_an_unsupported_layout(self) -> None:
        with self.assertRaisesRegex(ValueError, "Unsupported document text layout mode"):
            plan_pptx(_pptx_bytes(), None, "en", "fr", document_text_layout="unknown")


if __name__ == "__main__":
    unittest.main()
